#!/usr/bin/env python3
"""Generate the director-facing VTDT assessment report as a styled DOCX."""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "tmp/director-data.json"
OUTPUT_PATH = (
    Path(sys.argv[2])
    if len(sys.argv) > 2
    else ROOT / "docs/VTDT-karjeras-paligs-direktoram.docx"
)
ASSET_DIR = ROOT / "tmp/director-docx-assets"

PRIMARY = "292561"
PRIMARY_SOFT = "5F57A7"
ACCENT = "8F7CF3"
LILAC = "EEE9FF"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F7F4FF"
INK = "1F1D38"
MUTED = "5C5870"
WHITE = "FFFFFF"
GREEN = "27816B"
GOLD = "A05B11"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
BORDER = "D9D2EE"
FONT = "Calibri"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color, bold=False):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        set_style_font(style, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    set_style_font(caption, 9, MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    left = hp.add_run("VTDT karjeras palīgs")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    tab = hp.add_run("\tModeļa versija 2026.5")
    set_run_font(tab, size=8.5, color=MUTED)
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), alignment=2)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    label = fp.add_run("Lapa ")
    set_run_font(label, size=8.5, color=MUTED)
    add_page_field(fp)

    doc.core_properties.title = "VTDT karjeras un profesiju izvēles palīgs"
    doc.core_properties.subject = "Direktora informatīvais apraksts"
    doc.core_properties.author = "VTDT karjeras palīga projekta dokumentācija"
    doc.core_properties.keywords = "VTDT, karjera, profesijas, 9. klase, 10. klase"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != CONTENT_DXA:
        raise ValueError(f"Table widths must total {CONTENT_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def write_cell(cell, text, size=9, color=INK, bold=False, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths_dxa, font_size=9, header_size=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    header_size = header_size or font_size
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], HEADER_FILL)
        write_cell(
            table.rows[0].cells[index],
            header,
            size=header_size,
            color=PRIMARY,
            bold=True,
        )
    repeat_header(table.rows[0])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            if row_index % 2:
                shade_cell(cells[index], "FAFAFC")
            write_cell(cells[index], value, size=font_size)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    after.paragraph_format.line_spacing = 1.0
    return table


def add_hyperlink(paragraph, text, url, color=PRIMARY, underline=True):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    run_properties.append(color_element)
    if underline:
        underline_element = OxmlElement("w:u")
        underline_element.set(qn("w:val"), "single")
        run_properties.append(underline_element)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    run_properties.append(fonts)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_callout(doc, text, fill=LIGHT_FILL, color=PRIMARY):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=color)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge_name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:space"), "5")
        edge.set(qn("w:color"), ACCENT)
        borders.append(edge)
    p_pr.append(borders)
    return paragraph


def add_picture_with_alt(doc, path, width, alt_text, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def pil_font(size, bold=False):
    filename = "Arial Bold.ttf" if bold else "Arial.ttf"
    return ImageFont.truetype(f"/System/Library/Fonts/Supplemental/{filename}", size)


def centered_multiline(draw, box, text, font, fill):
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=4, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        ((left + right - width) / 2, (top + bottom - height) / 2 - bbox[1]),
        text,
        font=font,
        fill=fill,
        spacing=4,
        align="center",
    )


def draw_arrow(draw, points, color="#6F68A8", width=4):
    draw.line(points, fill=color, width=width, joint="curve")
    x, y = points[-1]
    previous_x, previous_y = points[-2]
    angle = math.atan2(y - previous_y, x - previous_x)
    length = 13
    spread = 0.55
    arrow = [
        (x, y),
        (x - length * math.cos(angle - spread), y - length * math.sin(angle - spread)),
        (x - length * math.cos(angle + spread), y - length * math.sin(angle + spread)),
    ]
    draw.polygon(arrow, fill=color)


def make_flow_diagram(path):
    image = Image.new("RGB", (1900, 700), "white")
    draw = ImageDraw.Draw(image)
    font = pil_font(25, bold=True)
    small = pil_font(20)
    nodes = {
        "start": (45, 280, 235, 390, "Sākums", PRIMARY),
        "mode": (290, 280, 510, 390, "Režīma\nizvēle", ACCENT),
        "quick": (590, 95, 830, 205, "10 jautājumi", PRIMARY_SOFT),
        "deep": (590, 465, 830, 575, "18 jautājumi", BLUE),
        "profile": (900, 280, 1150, 390, "22 dimensiju\nprofils", PRIMARY),
        "compare": (1230, 280, 1500, 390, "13 profesiju\nsalīdzinājums", PRIMARY_SOFT),
        "adaptive": (1580, 95, 1845, 205, "Adaptīvs\nprecizējums", GOLD),
        "result": (1580, 465, 1845, 575, "Top 3 / kopīga\n1. vieta", GREEN),
    }
    for left, top, right, bottom, label, color in nodes.values():
        draw.rounded_rectangle(
            (left, top, right, bottom),
            radius=20,
            fill=f"#{color}",
            outline="white",
            width=3,
        )
        centered_multiline(draw, (left, top, right, bottom), label, font, "white")

    draw_arrow(draw, [(235, 335), (290, 335)])
    draw_arrow(draw, [(510, 320), (550, 320), (550, 150), (590, 150)])
    draw_arrow(draw, [(510, 350), (550, 350), (550, 520), (590, 520)])
    draw_arrow(draw, [(830, 150), (865, 150), (865, 315), (900, 315)])
    draw_arrow(draw, [(830, 520), (865, 520), (865, 355), (900, 355)])
    draw_arrow(draw, [(1150, 335), (1230, 335)])
    draw_arrow(draw, [(1500, 315), (1540, 315), (1540, 150), (1580, 150)])
    draw_arrow(draw, [(1500, 355), (1540, 355), (1540, 520), (1580, 520)])
    draw_arrow(draw, [(1710, 95), (1710, 40), (1025, 40), (1025, 280)])
    draw.text((1600, 245), "ja vajag", font=small, fill=f"#{MUTED}")
    image.save(path)


def make_simulation_chart(path, professions):
    quick = {
        "apgerbu_dizainera_asistents": 12.85,
        "lauksaimniecibas_mehanizacijas_tehnikis": 17.35,
        "augkopibas_tehnikis": 8.67,
        "mebelu_galdnieks": 3.79,
        "apdares_darbu_tehnikis": 4.00,
        "eku_buvtehnikis": 8.10,
        "namdaris": 6.95,
        "arhitekturas_tehnikis": 5.34,
        "datorsistemu_tehnikis": 14.63,
        "programmesanas_tehnikis": 10.25,
        "automehanikis": 4.68,
        "autovirsbuvju_remonta_tehnikis": 2.09,
        "elektrotehnikis": 1.30,
    }
    deep = {
        "apgerbu_dizainera_asistents": 12.26,
        "lauksaimniecibas_mehanizacijas_tehnikis": 16.68,
        "augkopibas_tehnikis": 11.63,
        "mebelu_galdnieks": 2.96,
        "apdares_darbu_tehnikis": 1.74,
        "eku_buvtehnikis": 10.71,
        "namdaris": 5.93,
        "arhitekturas_tehnikis": 4.10,
        "datorsistemu_tehnikis": 12.76,
        "programmesanas_tehnikis": 8.16,
        "automehanikis": 2.90,
        "autovirsbuvju_remonta_tehnikis": 8.40,
        "elektrotehnikis": 1.77,
    }
    labels = [
        "Apģērbu\ndizains", "Lauks.\ntehnika", "Augkopība", "Mēbeles",
        "Apdare", "Ēkas", "Namdaris", "Arhitektūra", "Datorsist.",
        "Programm.", "Automeh.", "Virsbūves", "Elektro",
    ]
    ids = [profession["id"] for profession in professions]
    image = Image.new("RGB", (1900, 780), "white")
    draw = ImageDraw.Draw(image)
    font = pil_font(20)
    small = pil_font(15)
    bold = pil_font(21, bold=True)
    left, top, right, bottom = 125, 70, 1840, 625
    maximum = 21
    for tick in range(0, 21, 5):
        y = bottom - (tick / maximum) * (bottom - top)
        draw.line((left, y, right, y), fill="#DDD9E8", width=2)
        label = f"{tick}%"
        bbox = draw.textbbox((0, 0), label, font=font)
        draw.text((left - 18 - (bbox[2] - bbox[0]), y - 10), label, font=font, fill=f"#{MUTED}")
    threshold_y = bottom - (20 / maximum) * (bottom - top)
    draw.line((left, threshold_y, right, threshold_y), fill=f"#{GOLD}", width=4)
    group_width = (right - left) / len(ids)
    bar_width = group_width * 0.28
    for index, id_ in enumerate(ids):
        center = left + (index + 0.5) * group_width
        for offset, value, color in (
            (-bar_width, quick[id_], PRIMARY_SOFT),
            (0, deep[id_], BLUE),
        ):
            bar_left = center + offset
            bar_right = bar_left + bar_width
            bar_top = bottom - (value / maximum) * (bottom - top)
            draw.rectangle((bar_left, bar_top, bar_right, bottom), fill=f"#{color}")
        centered_multiline(
            draw,
            (center - group_width / 2, bottom + 16, center + group_width / 2, 745),
            labels[index],
            small,
            f"#{INK}",
        )
    draw.text((left, 18), "Top 1 īpatsvars 30 000 nejaušu profilu simulācijā", font=bold, fill=f"#{PRIMARY}")
    draw.rectangle((1240, 24, 1264, 48), fill=f"#{PRIMARY_SOFT}")
    draw.text((1274, 24), "Ātrais", font=small, fill=f"#{INK}")
    draw.rectangle((1370, 24, 1394, 48), fill=f"#{BLUE}")
    draw.text((1404, 24), "Padziļinātais", font=small, fill=f"#{INK}")
    draw.line((1580, 36, 1620, 36), fill=f"#{GOLD}", width=4)
    draw.text((1630, 24), "20% robeža", font=small, fill=f"#{INK}")
    image.save(path)


def add_cover(doc, source_url):
    for _ in range(4):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(18)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("DIREKTORA INFORMATĪVAIS APRAKSTS")
    set_run_font(run, size=10.5, color=ACCENT, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("VTDT karjeras un\nprofesiju izvēles palīgs")
    set_run_font(run, size=30, color=PRIMARY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run(
        "Divu režīmu, svērts un adaptīvs atbalsts 9.–10. klašu jauniešiem"
    )
    set_run_font(run, size=14, color=PRIMARY_SOFT)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(8)
    run = metadata.add_run("Modeļa versija 2026.5  ·  2026. gada 4. septembris")
    set_run_font(run, size=10.5, color=MUTED, bold=True)

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(70)
    add_hyperlink(source, "Oficiālais avots: VTDT profesiju katalogs", source_url)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run(
        "Karjeras izpētes orientieris — nevis zinātniski validēta psiholoģiska diagnoze"
    )
    set_run_font(run, size=10, color=MUTED, italic=True)
    doc.add_page_break()


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def matrix_rows(data, dimension_ids):
    return [
        [profession["title"]]
        + [f'{profession["profile"][dimension_id]:.2f}'.replace(".", ",") for dimension_id in dimension_ids]
        for profession in data["professions"]
    ]


def build_document(data):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    flow_path = ASSET_DIR / "user-flow.png"
    simulation_path = ASSET_DIR / "simulation.png"
    make_flow_diagram(flow_path)
    make_simulation_chart(simulation_path, data["professions"])

    doc = Document()
    configure_document(doc)
    add_cover(doc, data["sources"]["catalogue"])

    heading(doc, "Kopsavilkums vadībai", 1)
    add_callout(
        doc,
        "Rīks īsā, jaunietim saprotamā sarunā salīdzina atbilžu radīto 22 dimensiju profilu ar 13 aktuālajām VTDT profesijām. Viena atbilde profesiju nepiešķir, bet nepietiekams profils nesaņem noklusējuma līderi.",
    )
    add_table(
        doc,
        ["Rādītājs", "Vērtība", "Vadības nozīme"],
        [
            ["Režīmi", "Ātrais un padziļinātais", "Var izvēlēties situācijai atbilstošu ilgumu"],
            ["Pamata jautājumi", "10 vai 18", "Ātrais joprojām pārklāj visas 22 dimensijas"],
            ["Precizējumi", "līdz 3 vai 2", "Uzdod tikai tad, ja vajag papildu informāciju"],
            ["Profesijas", "13", "Tieši aktuālais VTDT katalogs"],
            ["Rezultāts", "Top 3 / kopīga 1. vieta", "Netiek mākslīgi paaugstināta pārliecība"],
        ],
        [2200, 2100, 5060],
        font_size=9.2,
    )

    heading(doc, "1. Mērķis un auditorija", 1)
    add_body(
        doc,
        "VTDT karjeras palīgs ir paredzēts galvenokārt 9.–10. klašu jauniešiem aptuveni 14–17 gadu vecumā. Tas nepaziņo vienu “pareizo” izvēli, bet palīdz pamanīt interešu, domāšanas veida un vēlamās darba vides kombināciju un piedāvā profesijas tālākai izpētei.",
    )
    add_body(
        doc,
        "Aplikācija ir statiska: nav backend, datubāzes vai ārēja vērtēšanas servisa. Aprēķins notiek lietotāja pārlūkā un publicēšanai GitHub Pages nav vajadzīgs build solis.",
    )

    heading(doc, "2. Ātrais un padziļinātais režīms", 1)
    add_table(
        doc,
        ["Kritērijs", "Ātrais tests", "Padziļinātais tests"],
        [
            ["Mērķis", "Ātra iepazīšanās un Top 3", "Plašāks salīdzinājums"],
            ["Pamata jautājumi", "10", "18"],
            ["Adaptīvi precizējumi", "0–3", "0–2"],
            ["Kopā", "10–13", "18–20"],
            ["Laiks", "ap 1–2 minūtēm", "ap 3 minūtēm"],
            ["Pārklājums", "visas 22 dimensijas", "visas 22 ar atkārtotiem mērījumiem"],
        ],
        [2400, 3480, 3480],
        font_size=9.2,
    )
    add_body(
        doc,
        "Ātrais tests ir vizuāli ieteiktais režīms. Plašākam salīdzinājumam jaunietis var turpināt padziļināti, saglabājot pirmo desmit jautājumu atbildes.",
    )
    add_picture_with_alt(
        doc,
        flow_path,
        6.45,
        "Plūsmas diagramma: režīma izvēle, 10 vai 18 jautājumi, dimensiju profils, profesiju salīdzinājums, iespējamais precizējums un Top 3.",
        "1. attēls. Lietotāja ceļš abos testa režīmos.",
    )

    heading(doc, "3. Aktuālās profesijas pa nozarēm", 1)
    sectors = {}
    for profession in data["professions"]:
        sectors.setdefault(profession["sector"], []).append(profession["title"])
    sector_rows = [
        [sector, "; ".join(titles)]
        for sector, titles in sorted(sectors.items(), key=lambda item: item[0])
    ]
    add_table(doc, ["Nozare", "Profesijas"], sector_rows, [2500, 6860], font_size=9.2)
    add_body(
        doc,
        "Inženiersistēmu būvtehniķis un Atjaunojamās enerģētikas tehniķis runtime datos nav iekļauti, jo tie nav aktuālajā VTDT profesiju katalogā. Vēsturiskais stāvoklis ir saglabāts projekta arhīva versijā.",
    )

    heading(doc, "4. Kāpēc jautājumi ir netieši", 1)
    add_body(
        doc,
        "Tiešs jautājums par programmēšanu, automašīnu remontu vai mēbeļu izgatavošanu ļautu viegli atminēt paredzēto profesiju. Tas vairāk mērītu jau zināmu amata nosaukuma simpātiju nekā darba stilu. Tāpēc jautājumi runā par kļūdas atrašanu, taustāmu rezultātu, telpisku iztēli, pacietību, mainīgiem apstākļiem, secību, sadarbību un citiem starpnozaru signāliem.",
    )
    add_callout(
        doc,
        "Katrs pamata jautājums ietekmē 3–6 dimensijas. Dati nesatur profesijas saņēmēju, un automatizēts vienas atbildes tests apliecina, ka viena izvēle nerada gala līderi.",
        fill=LILAC,
    )

    heading(doc, "5. Visi 18 pamata jautājumi un svari", 1)
    add_body(
        doc,
        "Tabulā norādīts dimensiju vektors pirms atbildes koeficienta. “Drīzāk jā” to reizina ar +1, “Jā” ar +2, bet “Nē” ar −2. Īsie dimensiju nosaukumi atbilst pilnajiem jēdzieniem pielikuma matricā.",
    )
    quick_ids = set(data["modes"]["quick"]["baseQuestionIds"])
    question_rows = []
    for question in data["questions"]:
        measured = "; ".join(
            f'{data["dimensions"][dimension_id]["shortLabel"]} ({str(weight).replace(".", ",")})'
            for dimension_id, weight in question["vector"].items()
        )
        question_rows.append(
            [
                question["number"],
                "Abi" if question["id"] in quick_ids else "Padziļ.",
                question["prompt"],
                measured,
            ]
        )
    add_table(
        doc,
        ["Nr.", "Režīms", "Jautājums", "Mērītās dimensijas (svars)"],
        question_rows,
        [520, 1150, 3370, 4320],
        font_size=8.0,
        header_size=8.2,
    )

    heading(doc, "6. Četru atbilžu koeficienti", 1)
    answer_rows = [
        [answer["label"], f'{answer["coefficient"]:+d}']
        for answer in data["answerScale"]
    ]
    add_table(
        doc,
        ["Atbilde", "Koeficients"],
        answer_rows,
        [5000, 4360],
        font_size=10,
    )
    add_body(
        doc,
        "Skalā nav neitrālas izvēles. Jaunietis izvēlas tuvāko no divām pozitīvām vai divām negatīvām atbildēm; neviena atbilde profesiju nepiešķir tieši.",
    )

    heading(doc, "7. Gala formula un piemērs", 1)
    group_rows = [
        [group["label"], len(data["dimensionsByGroup"][group_id]), f'{group["weight"] * 100:.0f}%']
        for group_id, group in data["dimensionGroups"].items()
    ]
    add_table(
        doc,
        ["Grupa", "Dimensijas", "Svars"],
        group_rows,
        [5500, 1900, 1960],
        font_size=9.5,
    )
    add_callout(
        doc,
        "atbilstības rādītājs = 100 × (0,25 × RIASEC + 0,55 × uzdevumi/intereses + 0,20 × vide/stils)",
        fill=HEADER_FILL,
    )
    add_body(
        doc,
        "Vienkāršots piemērs: ja trīs grupu sakritības ir 0,70, 0,80 un 0,60, tad 100 × (0,25 × 0,70 + 0,55 × 0,80 + 0,20 × 0,60) = 73,5. UI rāda 73,5 / 100, nevis 73,5% varbūtību.",
    )
    add_body(
        doc,
        "Katras dimensijas virzienu normalizē pret faktiski saturiski atbildētajiem jautājumiem. Uzticamība sasniedz pilnu spēku tikai pēc vairākiem saskanīgiem signāliem; pretrunīgas atbildes to samazina. Ātrā un padziļinātā rezultāti tādēļ paliek vienā 0–100 skalā.",
    )

    heading(doc, "8. Adaptīvo precizējumu princips", 1)
    add_body(
        doc,
        "Precizējumu var izraisīt ļoti maza Top 1–Top 2 starpība, praktiski vienāds Top 3 vai nepietiekams grupas pārklājums. Algoritms determinēti izvēlas neatbildēto jautājumu, kas vislabāk nošķir pašreizējo Top 2/Top 3, aizpilda pierādījumu trūkumu un nav pārāk līdzīgs jau uzdotajam. Kandidātu neuzdod, ja tā dimensijas jau ir pietiekami nosegtas.",
    )
    trigger_rows = [
        ["Top 1–Top 2 starpība", "< 0,40", "< 0,45"],
        ["Top 1–Top 3 diapazons", "< 0,80", "< 0,85"],
        ["Grupas pārklājums", "< 2 novērojumiem", "< 2 novērojumiem"],
        ["Maksimālais skaits", "3", "2"],
    ]
    add_table(
        doc,
        ["Iemesls / limits", "Ātrais", "Padziļinātais"],
        trigger_rows,
        [4300, 2530, 2530],
        font_size=9.5,
    )
    add_body(
        doc,
        "Ja pēc limita pirmo rezultātu starpība nepārsniedz 0,35, tie saglabā kopīgu pirmo vietu. Alfabētiskā secība netiek pasniegta kā uzvarētājs.",
    )
    heading(doc, "8.1. Precizējošo jautājumu banka", 2)
    clarifier_rows = []
    for index, question in enumerate(data["tieBreakers"], start=1):
        positive = sorted(
            ((dimension_id, weight) for dimension_id, weight in question["vector"].items() if weight > 0),
            key=lambda item: item[1],
            reverse=True,
        )[:3]
        negative = sorted(
            ((dimension_id, weight) for dimension_id, weight in question["vector"].items() if weight < 0),
            key=lambda item: item[1],
        )[:3]
        clarifier_rows.append(
            [
                index,
                question["prompt"],
                ", ".join(data["dimensions"][id_]["shortLabel"] for id_, _ in positive) or "—",
                ", ".join(data["dimensions"][id_]["shortLabel"] for id_, _ in negative) or "—",
            ]
        )
    add_table(
        doc,
        ["Nr.", "Jautājums", "“Jā” virziens", "Pretējais virziens"],
        clarifier_rows,
        [520, 4300, 2270, 2270],
        font_size=8.1,
    )

    heading(doc, "9. Validācija, skolēnu profili un simulācijas", 1)
    validation_rows = [
        ["Loģiskās personas", "13/13 padziļināti Top 1; 13/13 ātri Top 3"],
        ["10 jauniešu scenāriji", "Vizuāli radošam profilam apģērbu dizains abos režīmos ir Top 3"],
        ["Vienas atbildes tests", "nav gala līdera; Top 1–Top 2 ≤ 2,5"],
        ["Ātrais · 30 000", "Top 1: 1,30–17,35%; min Top 3: 11,19%; 10–13 jautājumi"],
        ["Padziļinātais · 30 000", "Top 1: 1,74–16,68%; min Top 3: 9,46%; 18–20 jautājumi"],
        ["Tukšs profils", "visām profesijām 50,0; leader = null"],
    ]
    add_table(doc, ["Pārbaude", "Rezultāts"], validation_rows, [3100, 6260], font_size=9.2)
    add_picture_with_alt(
        doc,
        simulation_path,
        6.45,
        "Stabiņu diagramma ar katras no 13 profesijām Top 1 īpatsvaru 30 000 profilu ātrajā un padziļinātajā simulācijā; visi stabiņi ir zem 20 procentiem.",
        "2. attēls. Nejaušo profilu Top 1 sadalījums. Simulācija pārbauda struktūru, nevis prognozē skolēnu izvēles.",
    )

    quick_ranks = {
        "eku_buvtehnikis": "2",
        "autovirsbuvju_remonta_tehnikis": "2",
    }
    persona_rows = []
    for profession in data["professions"]:
        quick_rank = quick_ranks.get(profession["id"], "1")
        persona_rows.append([profession["title"], quick_rank, "1"])
    heading(doc, "9.1. Personu rezultāti", 2)
    add_table(
        doc,
        ["Mērķa profesija", "Ātrais rangs", "Padziļinātais rangs"],
        persona_rows,
        [5600, 1880, 1880],
        font_size=8.8,
    )

    doc.add_page_break()
    heading(doc, "10. Profesiju un dimensiju svaru matrica", 1)
    add_body(
        doc,
        "Vērtības 0–1 ir redakcionāli, auditējami profesiju profili. Tie nav VTDT publicēti psihometriski normatīvi. Lai tabulas būtu salasāmas, 22 dimensijas sadalītas četrās daļās; saīsinājumu skaidrojums ir virs katras tabulas.",
    )

    matrix_groups = [
        ("10.1. RIASEC", data["dimensionsByGroup"]["riasec"]),
        ("10.2. Uzdevumi/intereses — 1. daļa", data["dimensionsByGroup"]["tasks"][:6]),
        ("10.3. Uzdevumi/intereses — 2. daļa", data["dimensionsByGroup"]["tasks"][6:]),
        ("10.4. Darba vide un stils", data["dimensionsByGroup"]["environment"]),
    ]
    for table_index, (title, dimension_ids) in enumerate(matrix_groups):
        if table_index:
            doc.add_page_break()
        heading(doc, title, 2)
        legend = " · ".join(
            f'{index + 1}: {data["dimensions"][dimension_id]["shortLabel"]}'
            for index, dimension_id in enumerate(dimension_ids)
        )
        add_body(doc, legend)
        headers = ["Profesija"] + [str(index + 1) for index in range(len(dimension_ids))]
        first_width = 3150
        remaining = CONTENT_DXA - first_width
        widths = [first_width]
        base_width = remaining // len(dimension_ids)
        widths.extend([base_width] * len(dimension_ids))
        widths[-1] += CONTENT_DXA - sum(widths)
        add_table(
            doc,
            headers,
            matrix_rows(data, dimension_ids),
            widths,
            font_size=7.5,
            header_size=8,
        )

    doc.add_page_break()
    heading(doc, "11. Ierobežojumi", 1)
    limitations = [
        ["Psihometrija", "Rīks nav zinātniski validēts psiholoģisks tests."],
        ["Profili", "Dimensiju svari ir redakcionāli pieņēmumi; VTDT avots apstiprina profesijas, nevis koeficientus."],
        ["Simulācija", "Nejaušas atbildes atrod strukturālas kļūdas, bet neatdarina īstus skolēnus."],
        ["Interpretācija", "Rezultāts nav uzņemšanas lēmums, spēju atzinums vai karjeras garantija."],
        ["Laiks", "Intereses mainās, un 10 vai 18 jautājumi nevar pilnībā aprakstīt profesijas ikdienu."],
    ]
    add_table(doc, ["Ierobežojums", "Ko tas nozīmē"], limitations, [2400, 6960], font_size=9.2)

    heading(doc, "12. Ieteikumi tālākai pārbaudei", 1)
    recommendations = [
        ["1", "Veikt pilotu ar vismaz 30–50 dažādu skolu 9.–10. klašu skolēniem."],
        ["2", "Īsā intervijā pārbaudīt saprotamību un to, vai atbildes virziens nav pārāk acīmredzams."],
        ["3", "Vieniem skolēniem salīdzināt ātrā un padziļinātā Top 3 stabilitāti."],
        ["4", "Iesaistīt VTDT programmu pedagogus profesiju profilu matricas ekspertu pārskatā."],
        ["5", "Rezultātus pārrunāt ar karjeras konsultantu, nesaucot to par psihometrisku validāciju bez atbilstošas metodikas."],
        ["6", "Mainīt svarus tikai pēc dokumentēta, atkārtojama pierādījuma un atkārtot visas simulācijas."],
    ]
    add_table(doc, ["Solis", "Ieteikums"], recommendations, [900, 8460], font_size=9.4)

    heading(doc, "13. Pārvaldība un atjaunināšana", 1)
    governance = [
        ["Pirms uzņemšanas cikla", "Pārbaudīt aktuālo VTDT profesiju katalogu."],
        ["Pēc datu maiņas", "Palaist npm run check un atjaunot Draw.io."],
        ["Pēc modeļa maiņas", "Palielināt assessmentVersion, lai droši ignorētu vecu localStorage stāvokli."],
        ["Pēc pilotēšanas", "Saglabāt izmaiņu pamatojumu, personas un simulāciju salīdzinājumu."],
    ]
    add_table(doc, ["Kad", "Darbība"], governance, [3000, 6360], font_size=9.2)
    source = doc.add_paragraph()
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(4)
    source_run = source.add_run("Avots: ")
    set_run_font(source_run, size=9, color=MUTED, bold=True)
    add_hyperlink(source, "VTDT profesiju katalogs", data["sources"]["catalogue"])

    return doc


def main():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document(data)
    document.save(OUTPUT_PATH)
    print(f"Izveidots: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
